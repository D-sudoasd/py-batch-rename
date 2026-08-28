"""Load a column of names from the first sheet of an Excel workbook."""

from __future__ import annotations

from pathlib import Path


def load_name_column(path: str | Path, column: int = 0) -> list[str]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            names: list[str] = []
            for row in ws.iter_rows(min_col=column + 1, max_col=column + 1, values_only=True):
                value = row[0]
                names.append("" if value is None else str(value).strip())
            return names
        finally:
            wb.close()

    if suffix == ".xls":
        raise ValueError("请将 .xls 另存为 .xlsx 后再导入")

    if suffix in {".csv", ".txt"}:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        names = []
        for line in text.splitlines():
            cell = line.split(",")[column] if "," in line else line
            names.append(cell.strip())
        return names

    raise ValueError(f"不支持的表格格式: {suffix or path.name}")
